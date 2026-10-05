import os
from pptx import Presentation
from src.extractors.base import BaseExtractor, ExtractedDocument, ExtractedPassage


def _walk(shapes):
    for sh in shapes:
        if sh.shape_type == 6 and hasattr(sh, "shapes"):  # group
            yield from _walk(sh.shapes)
        else:
            yield sh


class PPTXExtractor(BaseExtractor):
    def extract(self, file_path: str, raw_data_dir: str) -> ExtractedDocument:
        rel = self.rel_path(file_path, raw_data_dir)
        doc_id = self.generate_doc_id(file_path, raw_data_dir)
        try:
            prs = Presentation(file_path)
        except Exception as e:
            raise RuntimeError(f"PPTX open failed: {e}") from e
        doc = ExtractedDocument(doc_id=doc_id, filename=os.path.basename(file_path), file_path=file_path,
                                rel_path=rel, file_type="pptx")
        for n, slide in enumerate(prs.slides, start=1):
            parts, title = [], None
            for sh in _walk(slide.shapes):
                if getattr(sh, "has_table", False) and sh.has_table:
                    rows = [[c.text.strip() for c in r.cells] for r in sh.table.rows]
                    if len(rows) > 1:
                        hdr = rows[0]
                        for r in rows[1:]:
                            parts.append("[Table row] " + " | ".join(f"{h}: {v}" for h, v in zip(hdr, r) if v))
                elif sh.has_text_frame:
                    for para in sh.text_frame.paragraphs:
                        t = para.text.strip()
                        if t:
                            title = title or t
                            parts.append(t)
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
                nt = slide.notes_slide.notes_text_frame.text.strip()
                if nt:
                    parts.append(f"[Speaker notes] {nt}")
            if parts:
                doc.passages.append(ExtractedPassage(
                    passage_id=f"{doc_id}_s{n}", doc_id=doc_id, content="\n".join(parts), page_number=n,
                    section_title=f"Slide {n}: {title}"))
        return self.finalize(doc, "")
