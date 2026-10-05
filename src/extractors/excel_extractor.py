import datetime as dt
import os
from src.extractors.base import BaseExtractor, ExtractedDocument, ExtractedPassage


def _s(v) -> str:
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


class ExcelExtractor(BaseExtractor):
    def _rows(self, path, ext):
        if ext == ".xlsx":
            import openpyxl
            wb = openpyxl.load_workbook(path, data_only=True)
            for ws in wb.worksheets:
                yield ws.title, [[_s(c) for c in row] for row in ws.iter_rows(values_only=True)]
        elif ext == ".xls":
            import xlrd
            wb = xlrd.open_workbook(path)
            for sh in wb.sheets():
                yield sh.name, [[_s(c) for c in sh.row_values(i)] for i in range(sh.nrows)]
        else:
            raise RuntimeError(f"unsupported spreadsheet type {ext}")

    def extract(self, file_path: str, raw_data_dir: str) -> ExtractedDocument:
        rel = self.rel_path(file_path, raw_data_dir)
        doc_id = self.generate_doc_id(file_path, raw_data_dir)
        ext = os.path.splitext(file_path)[1].lower()
        doc = ExtractedDocument(doc_id=doc_id, filename=os.path.basename(file_path), file_path=file_path,
                                rel_path=rel, file_type=ext.lstrip("."))
        try:
            sheets = list(self._rows(file_path, ext))
        except Exception as e:
            raise RuntimeError(f"Spreadsheet open failed: {e}") from e
        idx = 0
        for sheet, rows in sheets:
            hdr_i = next((i for i, r in enumerate(rows) if sum(1 for c in r if c) >= 2), None)
            if hdr_i is None:
                continue
            headers = rows[hdr_i]
            for r_i in range(hdr_i + 1, len(rows)):
                vals = rows[r_i]
                if not any(vals):
                    continue
                idx += 1
                cells = [f"{headers[i] if i < len(headers) and headers[i] else f'Col{i+1}'}: {v}"
                         for i, v in enumerate(vals) if v]
                doc.passages.append(ExtractedPassage(
                    passage_id=f"{doc_id}_r{idx}", doc_id=doc_id,
                    content=f"[Sheet: {sheet} | Row {r_i+1}] " + " | ".join(cells),
                    section_title=f"Sheet: {sheet}", page_number=None, extraction_method="table",
                    metadata={"sheet": sheet, "row": r_i + 1, "row_cells": dict(zip(headers, vals))}))
        return self.finalize(doc, "")
